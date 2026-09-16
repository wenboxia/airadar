"""给黄金集 v2 做资料准备，并生成供人工标注的 Word 表。

用法：
  python3 tools/prep_golden_doc.py            # 准备资料（有缓存则跳过）+ 生成 Word
  python3 tools/prep_golden_doc.py --doc-only # 只用缓存重新生成 Word

产出：
  evals/golden_set/golden_draft.prep.json   AI 给每条写的资料（留档：标注者当时看到了什么）
  evals/golden_set/黄金集v2_标注表.docx      主人填写的表

AI 只做资料准备，不给收录建议（边界见 evals/golden_set/ASSIST_PROMPT.md）。
写资料的模型和负责打分的主力模型（DeepSeek）必须不是同一家——和 D10 给评测选裁判是同一个理由。
默认用裁判模型（Kimi）；它限流很紧、又是推理模型，实测约 2.4 分钟一条，
所以也可以换备用模型（GLM）并发跑：--provider fallback --workers 3。每条记下是谁写的。

依赖 python-docx，只在开发机上用，不进 requirements.txt。
"""
import argparse
import json
import os
import sqlite3
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from pipeline.config import load_config                    # noqa: E402
from pipeline.guards import Budget                          # noqa: E402
from pipeline.llm import LLMClient, LLMError, Provider      # noqa: E402
from pipeline.stages.classify import CATEGORIES            # noqa: E402

GS = os.path.join(ROOT, "evals", "golden_set")
DRAFT = os.path.join(GS, "golden_draft.jsonl")
CACHE = os.path.join(GS, "golden_draft.prep.json")
DOC = os.path.join(GS, "黄金集v2_标注表.docx")
DB_PATH = os.path.join(ROOT, "data", "knowledge.db")

TYPES = ["一手发布", "从业者观点", "公司动态", "开源项目", "论文", "教程", "快讯", "营销宣传"]

_SYS = f"""你在帮一个人准备资料，让他快速判断一条 AI 领域内容值不值得收进长期知识库。
判断由他本人做，你只负责把内容讲清楚。

输出 JSON：
{{"what": "2-3 句中文，这篇到底讲了什么、有什么具体结论或做法",
 "type": "从这里选一个：{' / '.join(TYPES)}",
 "three_months": "有用 / 一般 / 没用 ——这条信息三个月后还成立、还有参考价值吗",
 "three_months_why": "一句话说明上一项的依据"}}

限制：
- 不要给「建议收录 / 不建议收录」，也不要打分
- 只依据给你的原文，不补充原文没有的信息；原文太短就在 what 里直说「原文信息有限」
- 标题是元数据，不算原文内容"""

_DUP = """下面是编号的内容摘要。找出**在讲同一件事**的条目（同一次发布、同一个事件的不同报道）。
只报确定是同一件事的；主题相近但不是同一件事的不算。
输出 JSON：{"groups": [[编号, 编号, ...], ...]}，没有就输出 {"groups": []}"""


def _client(cfg, which: str):
    """返回 (客户端, 模型名)。只允许非主力的两家。"""
    conf = {"judge": (cfg.judge_base_url, cfg.judge_api_key, cfg.judge_model),
            "fallback": (cfg.fallback_base_url, cfg.fallback_api_key, cfg.fallback_model)}[which]
    if not all(conf):
        sys.exit(f"未配置 {which} 模型")
    if conf[2] == cfg.llm_model:
        sys.exit("写资料的模型不能和主力打分模型相同")
    return LLMClient([Provider(which, *conf)], Budget(600_000, 300), {}), conf[2]


def _load_draft():
    with open(DRAFT, encoding="utf-8") as f:
        return [json.loads(l) for l in f if l.strip() and not l.startswith("//")]


def prepare(rows, which: str, workers: int):
    cache = {}
    if os.path.exists(CACHE):
        with open(CACHE, encoding="utf-8") as f:
            cache = json.load(f)
    items = cache.setdefault("items", {})
    todo = [r for r in rows if r["url"] not in items]
    if not todo and cache.get("dup_groups") is not None:
        print("资料已全部缓存，跳过模型调用")
        return cache

    cfg = load_config()
    llm, model = _client(cfg, which)
    print(f"资料准备：{len(todo)} 条待写（模型 {model}，并发 {workers}）", flush=True)
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row

    texts = {}
    for r in todo:
        db = conn.execute("SELECT content, summary_long FROM items WHERE url=?",
                          (r["url"],)).fetchone()
        texts[r["url"]] = ((db and (db["content"] or db["summary_long"])) or "")[:3500]

    def _one(r):
        text = texts[r["url"]]
        out = llm.json_chat(_SYS, f"【标题（元数据）】{r['title']}\n----------\n"
                                  f"【原文】\n{text or '（无原文）'}", max_tokens=1200)
        if not out.get("what"):
            raise LLMError("empty")
        return {"by": model, "what": str(out["what"])[:300],
                "type": out.get("type") if out.get("type") in TYPES else "（未判定）",
                "three_months": str(out.get("three_months", ""))[:4],
                "three_months_why": str(out.get("three_months_why", ""))[:80]}

    # 逐条落盘（第一版全部跑完才写缓存，中途停掉一条都没留下）；写缓存只在主线程做
    fails = 0
    with ThreadPoolExecutor(max_workers=workers) as pool:
        futs = {pool.submit(_one, r): r for r in todo}
        for n, fut in enumerate(as_completed(futs), 1):
            r = futs[fut]
            try:
                items[r["url"]] = fut.result()
                _save(cache)
            except Exception as e:  # noqa: BLE001 单条失败重跑会补
                fails += 1
                print(f"  失败：{r['title'][:30]} {str(e)[:80]}", flush=True)
            if n % 10 == 0:
                print(f"  进度 {n}/{len(todo)}（失败 {fails}）", flush=True)

    cache["dup_groups"] = find_duplicates(rows, items, llm, conn, cache)
    _save(cache)
    missing = sum(1 for r in rows if r["url"] not in items)
    print(f"完成。缺资料 {missing} 条（重跑会补），重复组 {len(cache['dup_groups'] or [])} 组")
    return cache


def find_duplicates(rows, items, llm, conn, cache, size=20, step=10):
    """按发布时间排好序，20 条一窗、每次滑 10 条，各窗分别找「同一件事」，最后合并。

    不一次喂 100 条：推理模型会想很久，GLM 的服务端在约 65 秒时直接断开连接
    （异常是 ConnectionError(ProtocolError(None))）。同一事件的报道通常挤在一两天内，
    相邻窗口有一半重叠，按时间排序后不会漏掉。
    每个窗口的结果单独存盘，重跑只补失败的窗口（一次运行中途把备用模型的余额用完了，
    前面几个窗口的结果不该跟着作废）。有窗口没查完就返回 None，不把「没查完」当成「没有重复」。"""
    def when(r):
        row = conn.execute("SELECT published_at FROM items WHERE url=?", (r["url"],)).fetchone()
        return (row and row[0]) or ""
    order = sorted(range(len(rows)), key=lambda i: when(rows[i]))
    parent = list(range(len(rows) + 1))

    def root(x):
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    starts = list(range(0, max(len(order) - size, 0) + 1, step))
    if starts[-1] + size < len(order):
        starts.append(len(order) - size)
    done = cache.setdefault("dup_windows", {})
    complete = True
    for n, s in enumerate(starts, 1):
        window = [i + 1 for i in order[s:s + size]]          # 1 起的序号，和表格一致
        key = ",".join(map(str, window))
        if key not in done:
            lines = [f"{k}. {rows[k - 1]['title'][:80]} —— "
                     f"{items.get(rows[k - 1]['url'], {}).get('what', '')[:120]}" for k in window]
            try:
                out = llm.json_chat(_DUP, "\n".join(lines), max_tokens=4000, timeout=300)
            except LLMError as e:
                print(f"  重复检测第 {n}/{len(starts)} 窗失败：{str(e)[:100]}", flush=True)
                complete = False
                continue
            done[key] = [[k for k in g if isinstance(k, int) and k in window]
                         for g in out.get("groups", []) if isinstance(g, list)]
            _save(cache)
            print(f"  重复检测第 {n}/{len(starts)} 窗完成", flush=True)
        for g in done[key]:
            for k in g[1:]:
                parent[root(k)] = root(g[0])
    if not complete:
        return None
    groups = {}
    for k in range(1, len(rows) + 1):
        groups.setdefault(root(k), []).append(k)
    return [sorted(g) for g in groups.values() if len(g) > 1]


def _save(cache):
    with open(CACHE, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=1)


def _link(par, text, url):
    """python-docx 没有现成的超链接接口，只能手写 XML。"""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for tag, val in (("w:color", "1F5FBF"), ("w:u", "single")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        rpr.append(el)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    h.append(run)
    par._p.append(h)


def build_doc(rows, cache):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "PingFang SC"
    st.font.size = Pt(8.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")

    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.0))

    doc.add_heading("AIRadar 黄金集 v2 · 标注表（100 条）", level=1)
    intro = [
        "判断标准：三个月后，你还愿意在知识库里搜到它吗？",
        "只填最右边三列：【收录】填 y 或 n；【分类】收录的才填，1–3 个，用顿号分隔；【理由】一句话，面试被问到时就念这句。",
        "价值优先级：①一手发布 ②顶尖从业者的判断 ③影响技术格局的公司动态（收购、重大人事、算力、定价大改）"
        " ④开源项目（star 不是门槛）⑤论文（只收正在被实质讨论的）。营销稿、同一件事的重复报道、与 AI 无关的内容筛掉——哪怕来源权威。",
        "可选分类：" + "、".join(CATEGORIES),
        f"「讲了什么 / 类型 / 三个月后」由 AI（{'、'.join(sorted({v.get('by', '?') for v in cache['items'].values()}))}）"
        "根据原文整理，只是资料，不是建议；写资料的模型和给内容打分的 DeepSeek 不是同一家。"
        "「系统分/判断」是 pipeline 当时的原判（发布 ≥75、送审 50–75、丢弃 <50）。",
        "填完保存，然后跑：python3 evals/import_golden_docx.py",
    ]
    for line in intro:
        doc.add_paragraph(line)

    heads = ["序号", "标题（点击打开原文）", "来源·等级", "系统分/判断", "讲了什么", "类型",
             "三个月后", "重复提示", "【收录】y/n", "【分类】", "【理由】"]
    widths = [0.9, 4.6, 2.2, 1.6, 7.2, 1.5, 2.3, 1.4, 1.3, 2.2, 2.5]   # 合计 27.7cm = 页宽减页边距
    table = doc.add_table(rows=1, cols=len(heads))
    table.style = "Table Grid"
    table.autofit = False       # python-docx 里这就是固定列宽（写 w:tblLayout fixed）

    def _shade(cell, color):
        tc = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), color)
        tc.append(shd)

    hdr = table.rows[0]
    trpr = hdr._tr.get_or_add_trPr()
    rep = OxmlElement("w:tblHeader")        # 跨页时重复表头
    rep.set(qn("w:val"), "true")
    trpr.append(rep)
    for i, h in enumerate(heads):
        c = hdr.cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        _shade(c, "FFF2CC" if h.startswith("【") else "D9E2F3")

    dup_of = {}
    for g in cache.get("dup_groups") or []:
        for n in g:
            dup_of[n] = [x for x in g if x != n]

    route = {"published": "发布", "review": "送审", "discarded": "丢弃"}
    for i, r in enumerate(rows, 1):
        info = cache["items"].get(r["url"], {})
        pl = r.get("_pipeline", {})
        cells = table.add_row().cells
        cells[0].text = str(i)
        _link(cells[1].paragraphs[0], r["title"], r["url"])
        cells[2].text = f"{r['source']} · {r['tier']}"
        cells[3].text = f"{pl.get('score', '')} · {route.get(pl.get('auto_status'), '')}"
        cells[4].text = info.get("what", "（资料缺失，请直接看原文）")
        cells[5].text = info.get("type", "")
        tm = info.get("three_months", "")
        cells[6].text = f"{tm}：{info.get('three_months_why', '')}" if tm else ""
        cells[7].text = "同 " + "、".join(map(str, dup_of[i])) if i in dup_of else ""
        for j in (8, 9, 10):
            _shade(cells[j], "FFFBEB")

    for row in table.rows:
        for w, c in zip(widths, row.cells):
            c.width = Cm(w)
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(8.5)
    for p in doc.paragraphs[1:len(intro) + 1]:
        for run in p.runs:
            run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

    doc.save(DOC)
    print(f"已生成 {os.path.relpath(DOC, ROOT)}（{len(rows)} 行）")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--doc-only", action="store_true")
    ap.add_argument("--provider", choices=["judge", "fallback"], default="judge")
    ap.add_argument("--workers", type=int, default=1)
    args = ap.parse_args()
    rows = _load_draft()
    if args.doc_only:
        with open(CACHE, encoding="utf-8") as f:
            cache = json.load(f)
    else:
        cache = prepare(rows, args.provider, args.workers)
    build_doc(rows, cache)


if __name__ == "__main__":
    main()
